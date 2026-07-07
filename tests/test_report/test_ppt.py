"""Tests for ``mutadock.report.ppt_report``.

The deck is exercised against the synthetic ``sample_run_dir`` fixture (so the
suite runs anywhere) and, when present, against the real sample run. We assert
the slide *set* matches the data (structure/docking slides appear/vanish),
that pictures and tables are actually embedded, and that the deck round-trips
through ``save`` / re-open with a stable slide count.
"""

from __future__ import annotations

import io
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from mutadock.report import figures as F
from mutadock.report.data import RunData, discover_run
from mutadock.report.ppt_report import _b64_to_stream, build_pptx


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def _figures_for(run: RunData) -> dict:
    """Generate the four figures from a RunData (any may be None)."""
    structure = None
    if run.top_mutant_pdb is not None:
        sites = []
        if run.top_single_row is not None:
            try:
                sites = [
                    (
                        str(run.top_single_row["chain"]),
                        int(run.top_single_row["position"]),
                    )
                ]
            except (KeyError, ValueError, TypeError):
                sites = []
        structure = F.structure_backbone(run.top_mutant_pdb, sites)
    return {
        "ddg_distribution": F.ddg_distribution(run.single_df),
        "top_mutations": F.top_mutations_bar(run.single_df),
        "affinity": F.affinity_plot(run.docking_df),
        "structure": structure,
    }


def _all_shapes(prs: Presentation) -> list:
    return [shape for slide in prs.slides for shape in slide.shapes]


def _has_picture(prs: Presentation) -> bool:
    return any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in _all_shapes(prs))


def _has_table(prs: Presentation) -> bool:
    return any(s.shape_type == MSO_SHAPE_TYPE.TABLE for s in _all_shapes(prs))


def _all_text(prs: Presentation) -> str:
    chunks: list[str] = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                chunks.append(shape.text_frame.text)
            if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
                for row in shape.table.rows:
                    for cell in row.cells:
                        chunks.append(cell.text)
    return "\n".join(chunks)


# ---------------------------------------------------------------------------
# Core build (synthetic)
# ---------------------------------------------------------------------------
def test_build_pptx_synthetic(sample_run_dir: Path) -> None:
    run = discover_run(sample_run_dir)
    figs = _figures_for(run)
    prs = build_pptx(run, figs)

    # Duck-type Presentation (pptx.Presentation is a factory, not a class).
    assert hasattr(prs, "slides")
    # Title, summary, ddG, top-mut, docking, structure, provenance = 7.
    assert 5 <= len(prs.slides) <= 8

    # 16:9 canvas.
    from pptx.util import Inches

    assert prs.slide_width == Inches(13.333)
    assert prs.slide_height == Inches(7.5)

    assert _has_picture(prs)
    assert _has_table(prs)


def test_title_and_summary_text(sample_run_dir: Path) -> None:
    run = discover_run(sample_run_dir)
    figs = _figures_for(run)
    prs = build_pptx(run, figs)
    text = _all_text(prs)

    assert run.base_name in text  # "prot"
    stats = __import__(
        "mutadock.report.data", fromlist=["summary_stats"]
    ).summary_stats(run)
    # A concrete stat value should be present (mutations scored count).
    assert str(stats["n_scored"]) in text


def test_custom_title(sample_run_dir: Path) -> None:
    run = discover_run(sample_run_dir)
    prs = build_pptx(run, _figures_for(run), title="My Custom Deck")
    assert "My Custom Deck" in _all_text(prs)


# ---------------------------------------------------------------------------
# Save / reopen round-trip
# ---------------------------------------------------------------------------
def test_save_and_reopen(sample_run_dir: Path, tmp_path: Path) -> None:
    run = discover_run(sample_run_dir)
    prs = build_pptx(run, _figures_for(run))
    n = len(prs.slides)

    out = tmp_path / "r.pptx"
    prs.save(out)
    assert out.exists() and out.stat().st_size > 0

    reopened = Presentation(str(out))
    assert len(reopened.slides) == n


# ---------------------------------------------------------------------------
# Graceful degradation
# ---------------------------------------------------------------------------
def test_degradation_no_docking_no_structure(sample_run_dir: Path) -> None:
    run = discover_run(sample_run_dir)
    # Strip docking data.
    run.docking_df = None
    figs = _figures_for(run)
    figs["structure"] = None
    figs["affinity"] = None

    prs = build_pptx(run, figs)
    text = _all_text(prs)

    # No docking slide, no structure slide.
    assert "Docking affinity" not in text
    assert "Top mutant structure" not in text
    # Core slides survive.
    assert "Run summary" in text
    assert "ΔΔG distribution" in text
    assert run.base_name in text
    # Still valid / saveable.
    assert len(prs.slides) >= 4


def test_all_figures_none(sample_run_dir: Path) -> None:
    run = discover_run(sample_run_dir)
    figs = {
        "ddg_distribution": None,
        "top_mutations": None,
        "affinity": None,
        "structure": None,
    }
    prs = build_pptx(run, figs)
    text = _all_text(prs)
    # Title, summary, top-mutations (table only), docking (table only),
    # provenance still build; ddG image slide is skipped.
    assert "ΔΔG distribution" not in text
    assert "Run summary" in text
    # A table exists even with no figures (summary/provenance).
    assert _has_table(prs)


def test_minimal_rundata_no_crash() -> None:
    """A near-empty RunData still yields title + summary + provenance."""
    run = RunData(
        output_dir=Path("/tmp/none"),
        base_name="EMPTY",
        input_pdb=None,
        mutations_df=None,
        single_df=None,
        double_df=None,
        triple_df=None,
        docking_df=None,
        mutant_pdbs=[],
        top_mutant_pdb=None,
        top_single_row=None,
        metadata={
            "generated_at": "2026-07-02T00:00:00",
            "input": None,
            "output_dir": "/tmp/none",
            "files": [],
            "versions": {"pyrosetta": "x", "vina": "y"},
        },
    )
    prs = build_pptx(
        run,
        {
            "ddg_distribution": None,
            "top_mutations": None,
            "affinity": None,
            "structure": None,
        },
    )
    text = _all_text(prs)
    assert "EMPTY" in text
    assert "Run summary" in text
    assert "Provenance" in text
    assert len(prs.slides) == 3  # title, summary, provenance


# ---------------------------------------------------------------------------
# Helper unit
# ---------------------------------------------------------------------------
def test_b64_to_stream(sample_run_dir: Path) -> None:
    run = discover_run(sample_run_dir)
    b64 = F.ddg_distribution(run.single_df)
    assert b64 is not None
    stream = _b64_to_stream(b64)
    assert isinstance(stream, io.BytesIO)
    assert stream.tell() == 0
    assert stream.read(8) == b"\x89PNG\r\n\x1a\n"


# ---------------------------------------------------------------------------
# Real sample run
# ---------------------------------------------------------------------------
def test_build_pptx_real_sample(real_sample_dir: Path, tmp_path: Path) -> None:
    run = discover_run(real_sample_dir)
    figs = _figures_for(run)
    prs = build_pptx(run, figs)

    assert _has_picture(prs)
    assert _has_table(prs)
    assert 6 <= len(prs.slides) <= 8

    out = tmp_path / "real.pptx"
    prs.save(out)
    assert Presentation(str(out)).slides is not None
